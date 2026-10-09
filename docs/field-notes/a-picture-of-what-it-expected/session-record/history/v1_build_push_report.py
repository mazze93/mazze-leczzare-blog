"""Build blog-push-report_v1.0.0.xlsx from the 2026-10-08 push report.

Every row is something observed in this session (git output, GitHub's push
message) or explicitly marked unverified. No figures are computed in Python;
the summary counts are COUNTIF formulas so they follow Status edits.
"""
from pathlib import Path
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ---------------------------------------------------------------- content
REPO = "https://github.com/mazze93/mazze-leczzare-blog"
BRANCH = "field-note/a-picture-of-what-it-expected"

HEADERS = ["ID", "Item", "Category", "Status", "Observed evidence",
           "Verified?", "Not checked", "Next action", "Owner", "Link"]

ROWS = [
    ["P-01", f"Branch pushed: {BRANCH}", "Git", "Done",
     "git push printed '[new branch]' and set upstream tracking; 2 commits ahead of main (e165c22).",
     "Yes", "—", "Open a pull request when the draft is ready. Nothing merged yet.",
     "Mazze", f"{REPO}/pull/new/{BRANCH}"],
    ["P-02", "Push access to the blog repo added to this session", "Access", "Done",
     "First push got HTTP 403 from the session's git proxy (repo not authorized). Push succeeded after adding the repo with push access.",
     "Yes", "—", "None.", "Claude", REPO],
    ["P-03", "CSP fix commit hash corrected", "Record", "Corrected",
     "Earlier reply cited 83f7c60. git log shows 5f966ff; amending the commit message changed the hash.",
     "Yes", "—", "Use 5f966ff in any references.", "Claude", f"{REPO}/commit/5f966ff"],
    ["P-04", "Cloudflare Pages may build a preview of this branch", "Exposure", "Unverified",
     "draft: true hides the post from listings only. Files under public/ (images, logs, artifact page) would be served by any preview deployment.",
     "No", "Whether branch preview deployments are enabled for the Pages project.",
     "Check the Pages project's preview deployment settings.", "Mazze", ""],
    ["P-05", "Dependabot alerts on main: 4 (3 high, 1 moderate)", "Security", "Open",
     "Reported by GitHub in the push response. Not introduced by this branch.",
     "No", "Which packages, whether they reach production code, whether fixes exist.",
     "Review the alerts.", "Mazze", f"{REPO}/security/dependabot"],
]

COMMITS = [
    ["b8fbeda", "feat(blog)", 'draft field note "A Picture of What It Expected"'],
    ["5f966ff", "fix(csp)", "allow posts to frame their own /artifacts/ pages"],
]

STATUSES = ["Open", "Unverified", "Done", "Corrected"]

# ---------------------------------------------------------------- styles
FONT = "Arial"
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
MUTED = Font(name=FONT, size=9, italic=True, color="555555")
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
HEADER_FILL = PatternFill("solid", fgColor="1F2A44")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="C8CCD4")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

FIRST_ROW = 5          # first data row of the tracker
LAST_ROW = 50          # room for added rows; COUNTIFs and validation cover it
STATUS_COL = "D"


def write_header_row(ws, row, headers):
    """Write a dark header row. Side effect: styles cells in place."""
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font, cell.fill, cell.alignment, cell.border = HEADER_FONT, HEADER_FILL, WRAP, BOX


def write_body_row(ws, row, values, link_col=None):
    """Write one data row; the link column becomes a clickable hyperlink."""
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font, cell.alignment, cell.border = BODY, WRAP, BOX
        if col == link_col and value:
            cell.hyperlink, cell.font = value, LINK


def add_status_rules(ws):
    """Dropdown on Status plus fills keyed to it (whole row, so state reads at a glance)."""
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv.add(f"{STATUS_COL}{FIRST_ROW}:{STATUS_COL}{LAST_ROW}")
    ws.add_data_validation(dv)
    fills = {"Open": "FBE3E1", "Unverified": "FFF2CC", "Done": "E2F0D9", "Corrected": "DDEBF7"}
    span = f"A{FIRST_ROW}:J{LAST_ROW}"
    for status, color in fills.items():
        rule = FormulaRule(formula=[f'${STATUS_COL}{FIRST_ROW}="{status}"'],
                           fill=PatternFill("solid", fgColor=color))
        ws.conditional_formatting.add(span, rule)


def add_summary(ws, start_row):
    """COUNTIF summary + legend below the table. Returns nothing; writes cells."""
    rng = f"${STATUS_COL}${FIRST_ROW}:${STATUS_COL}${LAST_ROW}"
    ws.cell(row=start_row, column=1, value="Summary").font = BOLD
    for i, status in enumerate(STATUSES, start=1):
        ws.cell(row=start_row + i, column=1, value=status).font = BODY
        ws.cell(row=start_row + i, column=2, value=f'=COUNTIF({rng},"{status}")').font = BODY
    total_row = start_row + len(STATUSES) + 1
    ws.cell(row=total_row, column=1, value="Total").font = BOLD
    ws.cell(row=total_row, column=2,
            value=f"=SUM(B{start_row + 1}:B{total_row - 1})").font = BOLD

    legend = total_row + 2
    notes = [
        "How to use",
        "Edit Status (dropdown), Next action and Owner. Row colour follows Status.",
        "Observed evidence, Verified? and Not checked record what was seen on 2026-10-08; append a dated note rather than overwrite.",
        "Add new items below the last row (validation and counts cover rows 5–50).",
        "Source: Claude's push report, 2026-10-08 (git output and GitHub's push response). Nothing here was taken from the Cloudflare dashboard or the Dependabot page.",
    ]
    for i, text in enumerate(notes):
        cell = ws.cell(row=legend + i, column=1, value=text)
        cell.font = BOLD if i == 0 else MUTED


def build_tracker(wb):
    ws = wb.active
    ws.title = "Tracker"
    ws["A1"] = "Blog push report — field note branch"
    ws["A1"].font = TITLE
    ws["A2"] = f"Repo mazze93/mazze-leczzare-blog · branch {BRANCH} · 2026-10-08 · v1.0.0"
    ws["A2"].font = MUTED

    write_header_row(ws, FIRST_ROW - 1, HEADERS)
    for offset, row in enumerate(ROWS):
        write_body_row(ws, FIRST_ROW + offset, row, link_col=len(HEADERS))

    add_status_rules(ws)
    add_summary(ws, FIRST_ROW + len(ROWS) + 2)

    widths = [7, 34, 11, 12, 48, 10, 34, 34, 9, 44]
    for col, width in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = f"C{FIRST_ROW}"
    ws.auto_filter.ref = f"A{FIRST_ROW - 1}:J{FIRST_ROW + len(ROWS) - 1}"


def build_commits(wb):
    ws = wb.create_sheet("Commits")
    write_header_row(ws, 1, ["Hash", "Type", "Message", "Link"])
    for i, (sha, kind, msg) in enumerate(COMMITS, start=2):
        write_body_row(ws, i, [sha, kind, msg, f"{REPO}/commit/{sha}"], link_col=4)
    ws.cell(row=len(COMMITS) + 3, column=1,
            value="Base: main at e165c22. Both commits are on the branch only; neither is merged.").font = MUTED
    for col, width in zip("ABCD", [10, 11, 52, 60]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"


def main(out_path: Path) -> None:
    wb = Workbook()
    build_tracker(wb)
    build_commits(wb)
    wb.save(out_path)
    print(out_path)


if __name__ == "__main__":
    main(Path(sys.argv[1]))
